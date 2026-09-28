"""Independent G06/H10 rate precedence over every published operation.

Historical accepted rate is an input whose cohort correctness is audited under
H09, not inferred from passing this selection audit.
"""
from __future__ import annotations
import argparse,gzip,hashlib,json,os,re
from collections import Counter
from datetime import date,datetime,timedelta,timezone
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import load_workbook
from scripts.audit_planning_operation_hours import number,quantity,equal,BOOK,FOLDER,UNITS


def norm(value):return ' '.join(str(value or '').casefold().split()).replace('rectangular','retangular')

def operation_context(values,operation,primary,preparations):
 # Technical piece fields are common; each preparation supplies scheduling only.
 result=dict(values)
 matching=[p['values_json'] for p in preparations if str(p['values_json'].get('operation'))==operation]
 if matching:
  for field in ('machine','expected_date','planned_year','planned_week','operation'):
   if field in matching[-1]:result[field]=matching[-1][field]
 elif operation!=primary:
  result.update(machine=None,expected_date=None,planned_year=None,planned_week=None)
 return result

def rate_day(values,area,snapshot,periods,today):
 if values.get('expected_date'):return date.fromisoformat(str(values['expected_date'])[:10])
 if values.get('planned_year') is not None and values.get('planned_week') is not None:
  return date.fromisocalendar(int(values['planned_year']),int(values['planned_week']),1)
 if area=='cantoneiras' and values.get('imported_week') is not None:
  week=number(values['imported_week']);matches=[r for r in periods if r['area']==area and r['definition'].get('snapshot')==snapshot and r['definition'].get('week')==week]
  if len(matches)==1:return date.fromisocalendar(int(matches[0]['definition']['year']),int(week),1)
 return today

def choose(values,area,operation,resource_id,manual,history,excel,when):
 candidates=[]
 for record in manual:
  d=record['definition']
  if not d.get('confirmed') or (d.get('resource_id'),d.get('area'),str(d.get('operation')))!=(resource_id,area,operation):continue
  if not d['valid_from']<=when or d.get('valid_until') and when>d['valid_until']:continue
  if any(d.get(k) and norm(d[k])!=norm(values.get(k)) for k in ('material_type','profile')):continue
  candidates.append(record)
 if len(candidates)>1:return {'source':None,'rate':None,'candidates':sorted(str(r['id']) for r in candidates),'factor':1}
 if candidates:return {'source':'Manual','rate':candidates[0]['definition'],'configuration_id':str(candidates[0]['id']),'factor':1}
 h=number(history.get('value'))
 if h is not None and h>0:return {'source':'Histórico','rate':{k:history[k] for k in ('method','value','unit','window') if k in history},'factor':1}
 x=number((excel or {}).get('value'))
 if x is not None and x>0:
  factor=3 if (area,operation,values.get('machine'))==('perfis','corte','Serrote Fita Thomas IS639 Pav.1') and (quantity(values.get('quantity_required')) or 0)>50 else 1
  return {'source':'Excel provisório','rate':{**excel,'value':x*factor},'factor':factor}
 return {'source':None,'rate':None,'factor':1}

def main():
 parser=argparse.ArgumentParser();parser.add_argument('--output',default='c04-rate-selection');parser.add_argument('--workbook',type=Path,default=BOOK);args=parser.parse_args()
 assert re.fullmatch('[a-z0-9-]+',args.output)
 os.environ['MES_PG_DSN']=json.loads(Path('/home/luis/.local/state/planning-backups/integral-20260923/isolated.json').read_text())['dsn']
 from app import planning
 from app.raw import query
 from scripts.audit_planning_selected_ocr import source_fingerprints
 today=datetime.now(ZoneInfo(planning.settings.display_timezone)).date()
 bookhash=hashlib.sha256(args.workbook.read_bytes()).hexdigest();book=load_workbook(args.workbook,read_only=True,data_only=True)
 excel={str(r[1]):{'value':number(r[2]),'method':'area_hour','unit':'mm²/h'} for r in book['CapacidadeMáquinas'].iter_rows(min_row=2,max_row=12,max_col=3,values_only=True) if r[1]};book.close()
 report={'at':datetime.now(timezone.utc).isoformat(),'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),'rules':['G06','H10'],
  'environment':'planning_integral clone, repeatable read, no writes','day':str(today),'workbook_sha256':bookhash,'areas':{},'checks':0,'failures':[],
  'boundary':'Rate precedence verified independently from direct manual configurations, operation period and original Excel values. Historical cohort acceptance/exclusion remains H09. Real source ingestion remains C01.'}
 ledger=FOLDER/(args.output+'-rows.jsonl.gz')
 def check(name,actual,expected,**context):
  report['checks']+=1
  if not equal(actual,expected):report['failures'].append({'check':name,'actual':actual,'expected':expected,**context})
 with planning.connect(readonly=True) as c,gzip.open(ledger,'wt',encoding='utf8') as out:
  c.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ')
  assert c.execute('SELECT current_database() n').fetchone()['n']=='planning_integral'
  report['sources_before']=source_fingerprints(c)
  configs=[dict(r) for r in c.execute("SELECT * FROM planning_mtg.raw_objects WHERE NOT archived AND kind IN ('resource','rate','period') ORDER BY id")]
  report['configuration_sha256']=hashlib.sha256(json.dumps(configs,sort_keys=True,default=str).encode()).hexdigest()
  resources=[r for r in configs if r['kind']=='resource' and r['definition'].get('confirmed')]
  aliases={(a['area'],a['name']):r for r in resources for a in r['definition']['aliases']}
  manual=[r for r in configs if r['kind']=='rate'];periods=[r for r in configs if r['kind']=='period']
  history={r['hash']:r['proof'] for r in c.execute("SELECT hash,detail->'productivity' proof FROM planning_mtg.raw_contents WHERE detail ? 'productivity'")}
  hashes=set()
  for area in planning.AREAS:
   gen=query.generation(c,area);base,params=query.source(gen);snapshot=gen['metadata']['snapshot']['snapshot_id'];counts=Counter()
   rows=c.execute("SELECT m.row_key,c.values_json,c.detail->'calculation'->'operation_estimates' estimates,c.detail->'preparations' preparations,c.detail->'raw'->>'Mt\\h' excel_speed"+base,params)
   for row in rows:
    key=row['row_key'];v=row['values_json'];primary='corte' if area=='perfis' else str(v.get('operation') or '');counts['pieces']+=1
    for e in row['estimates'] or []:
     op=str(e['operation']);counts['operations']+=1;vals=operation_context(v,op,primary,row['preparations'] or [])
     when=rate_day(vals,area,snapshot,periods,today);resource=aliases.get((area,vals.get('machine')))
     resource_id=str(resource['id']) if resource else area+':'+str(vals.get('machine'))
     h=history.get(e.get('history_hash'));check('historical_evidence_exists',h is not None,True,area=area,key=key,operation=op)
     if h is None:continue
     hashes.add(e['history_hash'])
     days=resource['definition'].get('history_window_days',90) if resource else 90;end=min(today,when)
     window={'start':str(end-timedelta(days=days-1)),'end':str(end),'days':days}
     check('historical_window',h.get('window'),window,area=area,key=key,operation=op)
     method='area_hour' if area=='perfis' and op=='corte' else 'metres_hour' if area=='cantoneiras' else 'units_hour'
     check('historical_method',h.get('method'),method,area=area,key=key,operation=op)
     fallback=None
     if op==primary:
      fallback=excel.get(vals.get('machine')) if area=='perfis' else {'value':number(row['excel_speed']),'method':'metres_hour','unit':'m/h'}
     expected=choose(vals,area,op,resource_id,manual,h,fallback,str(when))
     allowed=(op in ('corte','abocardar') if area=='perfis' else op.isdigit() and op!='0') and (not resource or op in resource['definition']['operations'])
     if not allowed:expected.update(source=None,rate=None)
     check('chosen_source',e.get('source'),expected['source'],area=area,key=key,operation=op)
     check('chosen_factor',e.get('factor'),expected['factor'],area=area,key=key,operation=op)
     rate=expected['rate']
     for field in ('method','value'):
      check('chosen_rate_'+field,(e.get('rate') or {}).get(field),(rate or {}).get(field),area=area,key=key,operation=op)
     if expected.get('configuration_id') and allowed:check('manual_id',(e.get('configuration') or {}).get('id'),expected['configuration_id'],area=area,key=key,operation=op)
     if expected.get('candidates'):check('conflicting_candidates',sorted(e.get('candidates') or []),expected['candidates'],area=area,key=key,operation=op)
     if expected['source'] is None:check('unavailable_reason',bool(e.get('reason')) or e.get('quantity')==0,True,area=area,key=key,operation=op)
     if op==primary and area=='cantoneiras':
      check('G06_speed',v.get('speed_m_h'),rate.get('value') if rate and rate.get('method')=='metres_hour' else None,area=area,key=key)
      check('G06_unit',v.get('applied_rate_unit'),UNITS.get((rate or {}).get('method')),area=area,key=key)
     counts['source_'+str(expected['source'])]+=1
     counts['manual_conflict' if expected.get('candidates') else 'unambiguous']+=1
     out.write(json.dumps({'area':area,'key':key,'operation':op,'machine':vals.get('machine'),'resource_id':resource_id,'rate_date':str(when),
       'history_hash':e['history_hash'],'history_rate':h.get('value'),'historical_window':window,'excel_reference':fallback,
       'expected':expected,'observed':{'source':e.get('source'),'rate':e.get('rate'),'factor':e.get('factor'),'reason':e.get('reason')}},ensure_ascii=False,default=str)+'\n')
   report['areas'][area]={'generation':gen['id'],'counts':dict(counts)};print(area,dict(counts),'differences',len(report['failures']),flush=True)
  report['history_evidence_count']=len(hashes);report['sources_after']=source_fingerprints(c);check('sources_preserved',report['sources_after'],report['sources_before'])
 assert hashlib.sha256(args.workbook.read_bytes()).hexdigest()==bookhash
 report['ledger']={'file':ledger.name,'sha256':hashlib.sha256(ledger.read_bytes()).hexdigest()}
 report['result']='failed' if report['failures'] else 'passed';(FOLDER/(args.output+'.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+'\n')
 print(report['result'],report['checks'],'checks',len(report['failures']),'differences',flush=True)
 if report['failures']:raise SystemExit(1)

if __name__=='__main__':main()
