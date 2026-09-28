"""Independent F20/G07 operation-hour arithmetic and applied-rate provenance.

Reads current immutable published estimates and original configuration/workbook
rates. Does not call application estimation/rate-selection functions. Historical
cohort acceptance and priority of competing sources remain H09/H10 gates.
"""
from __future__ import annotations
import argparse,gzip,hashlib,json,math,os,re
from collections import Counter
from datetime import datetime,timezone
from pathlib import Path
from openpyxl import load_workbook

FOLDER=Path('docs/validacao-planeamento-integral/20260923-execucao')
BOOK=Path('/home/luis/projects/DATARESEARCHMTG/Met2_Plan_Perfis.xlsm')
UNITS={'area_hour':'mm²/h','metres_hour':'m/h','units_hour':'un./h','minutes_unit':'min/un.','fixed_minutes':'min'}

def number(value):
 if isinstance(value,bool):return None
 try:n=float(str(value).replace(',','.'));return n if math.isfinite(n) else None
 except (ValueError,TypeError):return None

def quantity(value):
 n=number(value);return n if n is not None and n>=0 and n.is_integer() else None

def equal(a,b):
 if b is None:return a is None
 if isinstance(b,(int,float)) and not isinstance(b,bool):
  n=number(a);return n is not None and abs(n-b)<=max(1e-6,abs(b)*1e-8)
 return a==b

def expected_hours(q,rate,length,section,allowed):
 if not allowed:return None
 if q is None:return None
 if q==0:return 0
 if not rate:return None
 speed=number(rate.get('value'));setup=number(rate.get('setup_minutes',0))
 if speed is None or speed<=0 or setup is None or setup<0:return None
 method=rate.get('method')
 if method=='area_hour':amount=q*section if section is not None and section>0 else None
 elif method=='metres_hour':amount=q*length/1000 if length is not None and length>0 else None
 elif method=='units_hour':amount=q
 elif method=='minutes_unit':return (q*speed+setup)/60
 elif method=='fixed_minutes':return (speed+setup)/60
 else:return None
 return amount/speed+setup/60 if amount is not None else None

def main():
 parser=argparse.ArgumentParser();parser.add_argument('--output',default='c04-operation-hours');args=parser.parse_args()
 assert re.fullmatch('[a-z0-9-]+',args.output)
 os.environ['MES_PG_DSN']=json.loads(Path('/home/luis/.local/state/planning-backups/integral-20260923/isolated.json').read_text())['dsn']
 from app import planning
 from app.raw import query
 from scripts.audit_planning_selected_ocr import source_fingerprints
 report={'at':datetime.now(timezone.utc).isoformat(),'script_sha256':hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
  'environment':'planning_integral clone, repeatable read, no writes','rules':['F20','G07'],'areas':{},'checks':0,'failures':[],
  'boundary':'Independent hours from selected counters and chosen rate. Direct manual/Excel rate provenance and stored historical rate checked; historical acceptance/exclusion and source priority are separate H09/H10 gates.'}
 before_book=hashlib.sha256(BOOK.read_bytes()).hexdigest();book=load_workbook(BOOK,read_only=True,data_only=True)
 excel={str(row[1]):{'value':number(row[2]),'cell':'C'+str(i)} for i,row in enumerate(book['CapacidadeMáquinas'].iter_rows(min_row=2,max_row=12,max_col=3,values_only=True),2) if row[1]};book.close()
 report['workbook']={'file':str(BOOK),'sha256':before_book,'rates':excel}
 ledger=FOLDER/(args.output+'-rows.jsonl.gz')
 def check(field,actual,expected,**context):
  report['checks']+=1
  if not equal(actual,expected):report['failures'].append({'field':field,'expected':expected,'observed':actual,**context})
 with planning.connect(readonly=True) as conn,gzip.open(ledger,'wt',encoding='utf8') as out:
  conn.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ')
  assert conn.execute('SELECT current_database() n').fetchone()['n']=='planning_integral'
  report['sources_before']=source_fingerprints(conn)
  configs={str(r['id']):dict(r) for r in conn.execute("SELECT * FROM planning_mtg.raw_objects WHERE NOT archived AND kind IN ('rate','resource')")}
  report['configuration_sha256']=hashlib.sha256(json.dumps(configs,sort_keys=True,default=str).encode()).hexdigest()
  resources=[r for r in configs.values() if r['kind']=='resource' and r['definition'].get('confirmed')]
  aliases={(a['area'],a['name']):r for r in resources for a in r['definition']['aliases']}
  resource_names={r['name']:r for r in resources}
  history={r['hash']:r['proof'] for r in conn.execute("SELECT hash,detail->'productivity' proof FROM planning_mtg.raw_contents WHERE detail ? 'productivity'")}
  for area in planning.AREAS:
   gen=query.generation(conn,area);base,params=query.source(gen)
   counts=Counter();examples={};items={}
   ig=query.generation(conn,area,dataset='capacity_items');ib,ip=query.source(ig)
   for item in conn.execute('SELECT m.row_key,c.values_json,c.detail'+ib,ip):items[(item['detail']['planning_key'],str(item['values_json']['operation']))]=item
   rows=conn.execute("SELECT m.row_key,c.values_json,c.detail->'calculation' calculation,c.detail->'raw'->>'Mt\\h' excel_speed"+base,params)
   for row in rows:
    key=row['row_key'];v=row['values_json'];cal=row['calculation'];estimates=cal.get('operation_estimates') or []
    sources={str(s['operation']):s for s in cal.get('production_sources') or []};q=quantity(v.get('quantity_required'))
    primary='corte' if area=='perfis' else str(v.get('operation') or '')
    counts['pieces']+=1;counts['active' if v.get('planning_active') else 'historical']+=1
    check('estimate_present',bool(estimates),True,area=area,key=key)
    check('estimate_operation_unique',len({str(e['operation']) for e in estimates}),len(estimates),area=area,key=key)
    for e in estimates:
     op=str(e['operation']);rate=e.get('rate');source=e.get('source');item=items.get((key,op));counts['operations']+=1
     source_row=sources.get(op);produced=quantity(source_row.get('value')) if source_row else None
     expected_q=max(q-produced,0) if q is not None and produced is not None else None
     if source_row is None:
      # Additional manually prepared operation: no imported/OCR source to infer.
      counts['without_production_source']+=1;expected_q=quantity(e.get('quantity'))
     check('balance_from_selected_counter',e.get('quantity'),expected_q,area=area,key=key,operation=op)
     resource=aliases.get((area,e.get('machine'))) or resource_names.get(e.get('machine'))
     allowed=(op in ('corte','abocardar') if area=='perfis' else op.isdigit() and op!='0') and (not resource or op in resource['definition']['operations'])
     length=number(v.get('length_mm'));section=number(v.get('section_unit'))
     expected=expected_hours(expected_q,rate,length,section,allowed)
     check('operation_hours',e.get('hours'),expected,area=area,key=key,operation=op)
     if expected is None:check('unavailable_reason',bool(e.get('reason')),True,area=area,key=key,operation=op)
     counts['hours_unknown' if expected is None else 'hours_zero' if expected==0 else 'hours_positive']+=1
     counts['rate_'+str(source)]+=1;counts['method_'+str((rate or {}).get('method'))]+=1
     factor=3 if source=='Excel provisório' and area=='perfis' and op=='corte' and v.get('machine')=='Serrote Fita Thomas IS639 Pav.1' and (q or 0)>50 else 1
     check('thomas_factor',e.get('factor'),factor,area=area,key=key,operation=op)
     if factor==3:counts['thomas_factor_3']+=1
     if source=='Manual':
      ident=(e.get('configuration') or {}).get('id');original=configs.get(ident)
      check('manual_configuration_exists',bool(original),True,area=area,key=key,operation=op)
      if original:
       check('manual_rate_definition',rate,original['definition'],area=area,key=key,operation=op)
       check('manual_rate_revision',(e.get('configuration') or {}).get('revision'),original['revision'],area=area,key=key,operation=op)
     elif source=='Histórico':
      proof=history.get(e.get('history_hash'))
      check('historical_evidence_exists',bool(proof),True,area=area,key=key,operation=op)
      if proof:
       for field in ('method','value','unit','window'):check('historical_'+field,(rate or {}).get(field),proof.get(field),area=area,key=key,operation=op)
     elif source=='Excel provisório':
      if area=='perfis':
       evidence=(rate or {}).get('source') or {};original=excel.get(v.get('machine'))
       check('excel_machine_reference',bool(original),True,area=area,key=key,operation=op)
       if original:
        check('excel_column_C_rate',rate.get('value'),original['value']*factor,area=area,key=key,operation=op)
        check('excel_cell',evidence.get('cell'),original['cell'],area=area,key=key,operation=op)
      else:check('excel_original_speed',rate.get('value'),number(row['excel_speed']),area=area,key=key,operation=op)
     else:check('no_source_no_rate',rate,None,area=area,key=key,operation=op)
     if op==primary:
      check('RAW_hours',v.get('theoretical_hours'),expected,area=area,key=key)
      check('RAW_source',v.get('rate_source'),source,area=area,key=key)
      check('RAW_rate',v.get('applied_rate_value'),(rate or {}).get('value'),area=area,key=key)
      check('RAW_unit',v.get('applied_rate_unit'),UNITS.get((rate or {}).get('method')),area=area,key=key)
      rule=(cal.get('rules') or {}).get('theoretical_hours') or {}
      check('hours_detail_unit',rule.get('unit'),'h',area=area,key=key)
      check('hours_detail_quantity',(rule.get('inputs') or {}).get('quantity'),expected_q,area=area,key=key)
      check('hours_detail_rate',(rule.get('inputs') or {}).get('rate'),rate,area=area,key=key)
      if area=='cantoneiras':check('speed_only_metres_per_hour',v.get('speed_m_h'),rate.get('value') if rate and rate.get('method')=='metres_hour' else None,area=area,key=key)
     if item:
      check('capacity_item_hours',item['values_json'].get('planned_hours'),expected,area=area,key=key,operation=op)
      counts['capacity_items_compared']+=1
     trace={'area':area,'key':key,'operation':op,'active':v.get('planning_active'),'quantity_required':q,'selected_counter':produced,'expected_balance':expected_q,
       'length_mm':length,'section_unit':section,'allowed_operation':allowed,'source':source,'rate':rate,'factor':factor,
       'expected_hours':expected,'observed_hours':e.get('hours'),'reason':e.get('reason'),'history_hash':e.get('history_hash')}
     out.write(json.dumps(trace,ensure_ascii=False,default=str)+'\n')
     sample=str(source)+':'+str((rate or {}).get('method'))+':'+('unknown' if expected is None else 'zero' if expected==0 else 'positive')
     examples.setdefault(sample,{'key':key,'operation':op,'expected':expected})
   report['areas'][area]={'generation':gen['id'],'capacity_generation':ig['id'],'counts':dict(counts),'examples':examples}
   print(area,dict(counts),'differences',len(report['failures']),flush=True)
  report['sources_after']=source_fingerprints(conn);check('sources_preserved',report['sources_after'],report['sources_before'])
 assert hashlib.sha256(BOOK.read_bytes()).hexdigest()==before_book
 report['ledger']={'file':ledger.name,'sha256':hashlib.sha256(ledger.read_bytes()).hexdigest()}
 report['result']='failed' if report['failures'] else 'passed'
 (FOLDER/(args.output+'.json')).write_text(json.dumps(report,ensure_ascii=False,indent=2,default=str)+'\n')
 print(report['result'],report['checks'],'checks',len(report['failures']),'differences',flush=True)
 if report['failures']:raise SystemExit(1)

if __name__=='__main__':main()
