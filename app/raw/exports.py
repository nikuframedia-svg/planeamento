"""Analytical downloads, independent of the operational XLSM writer."""
import csv,io,json,html
from datetime import date
from decimal import Decimal
from openpyxl import Workbook
from .. import planning,planning_needs as needs
from . import query,analysis,contracts


def safe_cell(v):
    if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')):return "'"+v
    return v


def table(p,format):
    if format not in ('csv','xlsx'):raise planning.PlanningError('Formato de exportação inválido.')
    with planning.connect(readonly=True) as c:
        c.execute('SET TRANSACTION ISOLATION LEVEL REPEATABLE READ')
        area=planning.check_area(p.get('area','perfis'));dataset=p.get('dataset','planning')
        gen=query.generation(c,area,p.get('version'),dataset);fields,exprs=query.field_expressions(c,area,dataset)
        columns=p.get('columns') or list(fields)
        if any(k not in fields for k in columns):raise planning.PlanningError('Coluna de exportação desconhecida.')
        base,params=query.source(gen);where,args=query.predicates(p,fields,exprs);base+=' AND '+where;params+=args
        selectargs=[];selects=[]
        for k in columns:
            e=exprs[k];selects.append(e.sql+' AS "'+k+'"');selectargs+=e.args
        order=p.get('order') or [];sorts=[];sortargs=[]
        for o in order:
            if o.get('field') not in fields or o.get('direction') not in ('asc','desc'):raise planning.PlanningError('Ordenação inválida.')
            e=exprs[o['field']];sorts.append('('+e.sql+') '+o['direction']+' NULLS LAST');sortargs+=e.args
        data=c.execute('SELECT '+','.join(selects)+base+' ORDER BY '+','.join(sorts+['m.row_key']),selectargs+params+sortargs).fetchall()
        rows=[{'values':r} for r in data];first={**gen['metadata'],'version':str(gen['id'])}
    labels=[safe_cell(fields[k]['label']) for k in columns]
    if format=='csv':
        buf=io.StringIO();writer=csv.writer(buf,delimiter=';');writer.writerow(labels)
        for r in rows:
            vals=[]
            for k in columns:
                v=r['values'].get(k)
                if isinstance(v,(int,float,Decimal)):
                    vals.append(str(v).replace('.',','));continue
                elif fields[k]['data_type']=='date' and v:
                    try:v=date.fromisoformat(str(v)).strftime('%d-%m-%Y')
                    except ValueError:pass
                vals.append(safe_cell(v))
            writer.writerow(vals)
        return '\ufeff'+buf.getvalue(),'text/csv; charset=utf-8'
    book=Workbook(write_only=True);sheet=book.create_sheet('RAW');sheet.append(labels)
    from openpyxl.cell import WriteOnlyCell
    for r in rows:
        cells=[]
        for k in columns:
            v=r['values'].get(k);cell=WriteOnlyCell(sheet,value=safe_cell(v))
            if fields[k]['data_type']=='date' and v:
                try:cell.value=date.fromisoformat(str(v));cell.number_format='DD-MM-YYYY'
                except ValueError:pass
            cells.append(cell)
        sheet.append(cells)
    source=book.create_sheet('Fontes');source.append(['Versão de consulta',first['version']]);source.append(['Fonte CPIS',first.get('cpis_mode')]);source.append(['População',p.get('population') or 'active']);source.append(['Consulta','Exportação analítica; valores desconhecidos permanecem vazios.'])
    b=io.BytesIO();book.save(b);return b.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'


def support_fields(result):
    """Freeze just the typed inputs of the saved definition, including formula dependencies."""
    selected={'of','ov','component_ref','machine','operation','production_date','status'}
    formulas={'calc_'+str(f['id']).replace('-',''):f['definition']['ast'] for f in result.get('formula_catalog',[])}
    def visit(node):
        if not isinstance(node,dict):return
        key=node.get('field')
        if key and key not in selected:
            selected.add(key)
            if key in formulas:visit(formulas[key])
        for child in node.get('args',[]):visit(child)
    for item in result['definition']['groups']+result['definition']['metrics']:visit(item['ast'])
    for item in result['filters'].get('filters',[]):visit(item)
    return sorted(selected)


def report(id,format='html'):
    job=analysis.get_job(id)
    if job['status']!='done' or job['kind']!='analysis':raise planning.PlanningError('Aguarda a conclusão da análise.',409)
    result=job['result'];payload={**job,'schema_version':2}
    selected=support_fields(result)
    with planning.connect(readonly=True) as c:
        hashes=list({r['hash'] for r in result['evidence']})
        # Keep source identifiers and calculation evidence, not complete unrelated Excel cells.
        payload['support']=c.execute("""SELECT hash,
            (SELECT jsonb_object_agg(k,v) FROM jsonb_each(values_json) t(k,v) WHERE k=ANY(%s)) AS values_json,
            detail - ARRAY['raw','original','preparations','sources','operations','cpis_conflicts'] AS detail,
            (SELECT jsonb_agg(s - 'payload') FROM jsonb_array_elements(coalesce(detail->'sources','[]')) s) AS sources
            FROM planning_mtg.raw_resolved_contents WHERE hash=ANY(%s) ORDER BY hash""",(selected,hashes)).fetchall()
    for row in payload['support']:
        calculation=row['detail'].get('calculation')
        if calculation:
            calculation['rules']={k:v for k,v in calculation.get('rules',{}).items() if k in selected}
    fields=contracts.mapping(result['definition']['area'],result['definition']['dataset'])
    payload['support_fields']={k:fields[k] for k in selected if k in fields}
    payload['support_scope']='Valores usados na definição e identificadores de suporte; versões e fórmulas conservadas no resultado.'
    if format=='json':return json.dumps(needs.serial(payload),ensure_ascii=False,indent=2),'application/json'
    if format!='html':raise planning.PlanningError('Formato inválido.')
    d=result['definition'];escape=lambda v:html.escape(str(v if v is not None else 'Por confirmar'))
    headings=[('g'+str(i),'Agrupamento '+str(i+1)) for i,_ in enumerate(d['groups'])]+[('m'+str(i),m['name']+((' ('+m['unit']+')') if m.get('unit') else '')) for i,m in enumerate(d['metrics'])]
    thead='<tr>'+''.join('<th>'+escape(label)+'</th>' for k,label in headings)+'<th>Cobertura</th></tr>'
    body=''
    for r in result['groups']:
        body+='<tr>'+''.join('<td>'+escape(r.get(k))+'</td>' for k,_ in headings)+'<td>'+escape(', '.join(f"{r['known'+str(i)]}/{r['rows']}" for i in range(len(d['metrics']))))+'</td></tr>'
    data=json.dumps(needs.serial(payload),ensure_ascii=False).replace('<','\\u003c').replace('&','\\u0026')
    rendered='''<!doctype html><html lang="pt"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>'''+escape(d.get('title','Análise RAW'))+'''</title><style>body{font:16px system-ui;color:#243b43;margin:24px;max-width:1200px}h1{font-size:24px}table{border-collapse:collapse;width:100%;font-size:14px}td,th{padding:8px;border-bottom:1px solid #dce3e6;text-align:left}th{position:sticky;top:0;background:#eef4f5}input,button{font:inherit;padding:7px;margin:8px}svg{max-width:100%;height:auto}.scroll{max-height:60vh;overflow:auto}.muted{color:#52636b}pre{white-space:pre-wrap}</style><h1>'''+escape(d.get('title','Análise RAW'))+'''</h1><p>Exportação fixa · '''+escape(result['created_at'])+''' · '''+escape(result['rows'])+''' linhas abrangidas</p><p class="muted">'''+escape(' '.join(result['limitations']))+'''</p><p>Interpretação do modelo: '''+escape(d.get('explanation',''))+'''</p><div id="chart"></div><input id="search" aria-label="Pesquisar resultados" placeholder="Pesquisar resultados"><button id="download">Guardar JSON</button><div class="scroll"><table><thead>'''+thead+'''</thead><tbody>'''+body+'''</tbody></table></div><details><summary>Definição, fórmulas e fontes</summary><pre>'''+escape(json.dumps({k:result[k] for k in ('definition','sources','filters')},ensure_ascii=False,indent=2))+'''</pre></details><script id="data" type="application/json">'''+data+'''</script><script>
const payload=JSON.parse(document.getElementById('data').textContent);document.getElementById('search').oninput=e=>{for(const r of document.querySelectorAll('tbody tr'))r.hidden=!r.textContent.toLocaleLowerCase('pt-PT').includes(e.target.value.toLocaleLowerCase('pt-PT'))};document.getElementById('download').onclick=()=>{const a=document.createElement('a');a.href=URL.createObjectURL(new Blob([JSON.stringify(payload,null,2)],{type:'application/json'}));a.download='analise-raw.json';a.click();URL.revokeObjectURL(a.href)};
RawCharts.render(document.getElementById('chart'),payload.result,i=>{const r=document.querySelectorAll('tbody tr')[i];if(r){r.scrollIntoView({block:'center'});r.style.background='#e3f2f4'}});
</script></html>'''
    from pathlib import Path
    chart_code=(Path(__file__).parents[1]/'web/static/raw_charts.js').read_text()
    rendered=rendered.replace('<script>\nconst payload=', '<script>'+chart_code+'\nconst payload=')
    return rendered,'text/html; charset=utf-8'
