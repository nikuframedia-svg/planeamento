"""Read-only workbook evidence, pinned to the exact imported file hash."""
from __future__ import annotations
import hashlib, json, os, subprocess, re
from pathlib import Path
from contextlib import contextmanager
from contextvars import ContextVar
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from datetime import datetime, timezone, timedelta
from psycopg.types.json import Jsonb
from openpyxl.formula.translate import Translator
from .. import planning, planning_needs as needs

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
SHEETS = {'perfis': {'CapacidadeMáquinas', 'PlanDisponibilidadeSemanal'},
          'cantoneiras': {'Plan_semanal', 'Dados', 'Analise maq'}}


# Formatos de data internos do Excel (numFmtId) e símbolos de data num formato próprio.
DATE_FORMATS = set(range(14, 23)) | {45, 46, 47}
DATE_TOKENS = re.compile(r'[dmyhs]', re.I)


def date_styles(z):
    """Índices de estilo (atributo s da célula) que o Excel mostra como data."""
    if 'xl/styles.xml' not in z.namelist():
        return set()
    root = ET.fromstring(z.read('xl/styles.xml'))
    custom = set()
    for fmt in root.iter(NS+'numFmt'):
        # Texto entre aspas, [cor]/[$-locale] e escapes não são símbolos de data.
        code = re.sub(r'"[^"]*"|\[[^\]]*\]|\\.', '', fmt.get('formatCode') or '')
        if DATE_TOKENS.search(code):
            custom.add(int(fmt.get('numFmtId')))
    xfs = root.find(NS+'cellXfs')
    return {i for i, xf in enumerate(xfs if xfs is not None else []) if int(xf.get('numFmtId') or 0) in DATE_FORMATS | custom}


def excel_date(value):
    """Número de série do Excel (sistema 1900) como data ISO; o valor guardado não muda."""
    if not isinstance(value, (int, float)) or not 1 <= value < 2958466:
        return None
    moment = datetime(1899, 12, 30) + timedelta(days=value)
    return moment.date().isoformat() if moment.time() == datetime.min.time() else moment.isoformat(timespec='minutes')


def extract(path, area):
    result = {}
    with ZipFile(path) as z:
        dates = date_styles(z)
        strings = []
        if 'xl/sharedStrings.xml' in z.namelist():
            with z.open('xl/sharedStrings.xml') as f:
                for _, el in ET.iterparse(f, events=('end',)):
                    if el.tag == NS+'si':
                        strings.append(''.join(el.itertext())); el.clear()
        relations = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
        workbook = ET.fromstring(z.read('xl/workbook.xml'))
        for sheet in workbook.find(NS+'sheets'):
            name = sheet.get('name')
            if name not in SHEETS[area]:
                continue
            target = relations[sheet.get(REL+'id')]
            target = target.lstrip('/') if target.startswith('/') else 'xl/'+target
            output = []; shared = {}
            with z.open(target) as f:
                for _, el in ET.iterparse(f, events=('end',)):
                    if el.tag != NS+'row' or not el.get('r'):
                        continue
                    cells = {}
                    for cell in el:
                        v = cell.find(NS+'v'); formula = cell.find(NS+'f')
                        value = v.text if v is not None else None
                        if cell.get('t') == 's' and value is not None:
                            value = strings[int(value)]
                        elif cell.get('t') == 'inlineStr':
                            value = ''.join(cell.find(NS+'is').itertext())
                        elif value is not None and cell.get('t') not in ('str', 'e'):
                            try:
                                value = float(value)
                                if value.is_integer(): value = int(value)
                            except ValueError: pass
                        text = formula.text if formula is not None else None
                        if formula is not None and formula.get('t') == 'shared':
                            if text: shared[formula.get('si')] = (text, cell.get('r'))
                            elif formula.get('si') in shared:
                                source, origin = shared[formula.get('si')]
                                text = Translator('='+source, origin=origin).translate_formula(cell.get('r'))[1:]
                        if value is not None or text:
                            item = {'cell': cell.get('r'), 'value': value, 'formula': text,
                                    'kind': 'formula' if text else 'error' if cell.get('t') == 'e' else 'value'}
                            # Auditoria 06/10 (C2-6): o Excel guarda as datas como número de série (46231);
                            # mantém-se o valor e junta-se a data legível quando a célula tem formato de data.
                            shown = excel_date(value) if cell.get('s') and int(cell.get('s')) in dates else None
                            if shown:
                                item['date'] = shown
                            cells[re.sub(r'\d', '', cell.get('r'))] = item
                    if cells: output.append({'row': int(el.get('r')), 'cells': cells})
                    el.clear()
            result[name] = output
    return result


_interactive = ContextVar('workbooks_interactive', default=False)


@contextmanager
def interactive():
    """Gravações de quem está no ecrã (registo, pré-visualização, Tabela): sem 409 por causa do Excel no disco."""
    token = _interactive.set(True)
    try:
        yield
    finally:
        _interactive.reset(token)


# Pastas que a carga automática (sync_drive.sh) lê, a raiz primeiro: desde 08/10 carrega o plano
# mais recente, pelo nome exato, entre a raiz e SAIDA/. MES_RAW_WORKBOOK_FOLDERS muda-as («.,SAIDA»).
AUTO_FOLDERS = ('', 'SAIDA')


def workbook_root():
    return Path(os.getenv('MES_RAW_WORKBOOK_ROOT', '/home/luis/projects/DATARESEARCHMTG'))


def auto_folders():
    raw = os.getenv('MES_RAW_WORKBOOK_FOLDERS')
    if raw is None: return AUTO_FOLDERS
    return tuple(dict.fromkeys('' if f.strip() in ('', '.') else f.strip().strip('/') for f in raw.split(',')))


def _sha256(path):
    digest = hashlib.sha256()
    with open(path, 'rb') as f:
        for block in iter(lambda: f.read(1 << 20), b''):
            digest.update(block)
    return digest.hexdigest()


def imported_file(src, root=None):
    """Caminho relativo do Excel importado na pasta local (ex.: 'SAIDA/Met3_Plan_Cantoneiras.xlsm').

    Desde 08/10 o loader guarda-o em source_filename; os snapshots anteriores só têm o nome,
    e a pasta sai do source_path quando este fica dentro da pasta local.
    """
    name = str(src.get('source_filename') or '')
    if '/' in name: return name
    try:
        return Path(src['source_path']).relative_to(root or workbook_root()).as_posix()
    except (KeyError, TypeError, ValueError):
        return name or None


def local_copies(src, root=None, folders=None):
    """Cópias locais do plano importado que podem ter o sha dele: a que o registo indica, depois raiz e SAIDA/."""
    root = Path(root or workbook_root())
    name = Path(str(src.get('source_filename') or '')).name
    rels = [imported_file(src, root)] + ([f'{f}/{name}' if f else name for f in (auto_folders() if folders is None else folders)] if name else [])
    return [root / rel for rel in dict.fromkeys(r for r in rels if r)]


def locate(src, root=None, folders=None):
    """(ficheiro com o sha importado, havia cópias) — o plano pode ter vindo da raiz ou de SAIDA/ (08/10)."""
    existing = [p for p in local_copies(src, root, folders) if p.is_file()]
    return next((p for p in existing if _sha256(p) == src['source_sha256']), None), bool(existing)


def capture(conn, area, path=None):
    snap = planning.snapshot(conn, area)
    if conn.execute('SELECT 1 FROM planning_mtg.raw_workbook_evidence WHERE snapshot_id=%s', (snap['snapshot_id'],)).fetchone():
        return
    src = conn.execute('SELECT source_sha256,source_filename,source_path FROM audit_mtg.snapshots WHERE snapshot_id=%s', (snap['snapshot_id'],)).fetchone()
    path = Path(path) if path else None
    if path and path.is_file():
        found, existing = (path if _sha256(path) == src['source_sha256'] else None), True
    else:
        # Procura entre as cópias locais conhecidas a que tem o sha importado, em vez de assumir a raiz (08/10).
        found, existing = locate(src, path)
    if not existing: return
    if not found:
        # Quem grava no formulário ou na Tabela não recebe 409 (07/10/2026): valem os valores importados.
        # O worker continua a parar, para não publicar horas de turno vazias.
        if _interactive.get(): return
        raise planning.PlanningError('A cópia Excel difere da versão importada; evidência de fórmulas pendente.', 409)
    sheets = extract(found, area)
    if _sha256(found) != src['source_sha256']:
        if _interactive.get(): return
        raise planning.PlanningError('O Excel mudou durante a leitura. Repete a consulta.', 409)
    conn.execute('INSERT INTO planning_mtg.raw_workbook_evidence(snapshot_id,source_sha256,source_filename,sheets) VALUES(%s,%s,%s,%s) ON CONFLICT DO NOTHING',
                 (snap['snapshot_id'], src['source_sha256'], src['source_filename'], Jsonb(sheets)))


def source(conn, area):
    snap = planning.snapshot(conn, area)
    row = conn.execute('SELECT * FROM planning_mtg.raw_workbook_evidence WHERE snapshot_id=%s', (snap['snapshot_id'],)).fetchone()
    if row: return needs.serial(row)
    # Cached values remain inspectable if the original XLSM is unavailable.
    sheets = {}
    from openpyxl.utils.cell import get_column_letter
    for r in conn.execute('SELECT sheet_name,excel_row,row_data FROM raw_mtg.other_sheet_rows WHERE snapshot_id=%s ORDER BY sheet_name,excel_row', (snap['snapshot_id'],)).fetchall():
        if r['sheet_name'] not in SHEETS[area]: continue
        cells = {get_column_letter(i): {'cell':get_column_letter(i)+str(r['excel_row']), 'value':v, 'formula':None, 'kind':'cached_value'}
                 for i,v in enumerate(r['row_data'].get('values', []), 1) if v is not None}
        sheets.setdefault(r['sheet_name'], []).append({'row':r['excel_row'], 'cells':cells})
    return {**needs.serial(snap), 'sheets':sheets, 'formula_status':'Original da versão indisponível; fórmulas por verificar.'}


DRIVE_EXCLUDES = ('crm-backups/**', 'SAIDA/backups/**')


def _drive_time(value):
    """ModTime do rclone ('…T07:04:10.123456789Z') como datetime comparável."""
    text = re.sub(r'(\.\d{6})\d+', r'\1', str(value or '')).replace('Z', '+00:00')
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        return datetime.min.replace(tzinfo=timezone.utc)


def drive_choice(entries, name):
    """O ficheiro a comparar com a importação: o que a carga automática lê ou um mais recente noutra subpasta.

    Auditoria 06/10 (A1-F1/A1-F2): o planeador já gravou o plano só em
    "Kanban's MTG3/" e em "SAIDA/". Desde 08/10 a carga automática lê o mais
    recente, pelo nome exato, entre a raiz e SAIDA/ (empate ou conteúdo igual:
    a raiz). Uma cópia noutra subpasta só conta quando é mais recente e
    diferente dessa; assim a app avisa que há versão por importar e diz onde está.
    """
    found = [r for r in entries if r.get('Name') == name and r.get('Hashes', {}).get('sha256')]
    auto = None
    for folder in auto_folders():
        r = next((r for r in found if r.get('Path', name) == (f'{folder}/{name}' if folder else name)), None)
        if r and (auto is None or (_drive_time(r.get('ModTime')) > _drive_time(auto.get('ModTime'))
                                   and r['Hashes']['sha256'] != auto['Hashes']['sha256'])):
            auto = r
    newest = max(found, key=lambda r: _drive_time(r.get('ModTime')), default=None)
    if auto and newest and newest is not auto and (newest['Hashes']['sha256'] == auto['Hashes']['sha256']
                                                  or _drive_time(newest.get('ModTime')) <= _drive_time(auto.get('ModTime'))):
        return auto
    return newest


def observe_drive(force=False):
    """Metadata only; never runs the importer or changes its twice-daily schedule."""
    if os.getenv('MES_RAW_DRIVE_CHECK', '0') != '1': return
    with planning.connect() as c:
        if not force and c.execute("SELECT 1 FROM planning_mtg.raw_drive_observations WHERE attempted_at>now()-interval '5 minutes' LIMIT 1").fetchone(): return
        for area in planning.AREAS:
            c.execute('INSERT INTO planning_mtg.raw_drive_observations(area) VALUES(%s) ON CONFLICT(area) DO UPDATE SET attempted_at=now()', (area,))
    names = [('perfis','Met2_Plan_Perfis.xlsm'),('cantoneiras','Met3_Plan_Cantoneiras.xlsm')]
    try:
        # Todas as pastas, só os dois ficheiros de planeamento (os backups do CRM ficam de fora).
        command = [os.getenv('MES_RCLONE_BIN', '/home/luis/bin/rclone'), 'lsjson', 'gdrive:', '--files-only', '--hash', '--recursive']
        # --filter (não --include/--exclude misturados): ordem garantida e pastas excluídas nem são listadas.
        for pattern in DRIVE_EXCLUDES: command += ['--filter', '- '+pattern]
        for _, name in names: command += ['--filter', '+ '+name]
        command += ['--filter', '- **']
        run = subprocess.run(command, capture_output=True, text=True, timeout=40, check=True)
        entries = json.loads(run.stdout)
        with planning.connect() as c:
            for area, name in names:
                r = drive_choice(entries, name)
                if not r: raise ValueError('missing hash')
                # remote_filename guarda o caminho; fora da raiz e de SAIDA/ a carga automática não o lê.
                c.execute('UPDATE planning_mtg.raw_drive_observations SET checked_at=now(),remote_sha256=%s,remote_modified_at=%s,remote_filename=%s,error=NULL WHERE area=%s',
                          (r['Hashes']['sha256'],r['ModTime'],r.get('Path', name),area))
    except Exception as exc:
        with planning.connect() as c:
            c.execute('UPDATE planning_mtg.raw_drive_observations SET error=%s', ('Não foi possível verificar o Drive ('+type(exc).__name__+'). Última informação conservada.',))


def status():
    with planning.connect(readonly=True) as c:
        output = []
        for area in planning.AREAS:
            snap = planning.snapshot(c, area)
            source = c.execute('SELECT source_sha256,loaded_at,source_filename,source_path FROM audit_mtg.snapshots WHERE snapshot_id=%s',(snap['snapshot_id'],)).fetchone()
            imported = {k:source[k] for k in ('source_sha256','loaded_at')}
            observed = c.execute('SELECT * FROM planning_mtg.raw_drive_observations WHERE area=%s',(area,)).fetchone()
            gen = c.execute('SELECT id,created_at,metadata FROM planning_mtg.raw_generations WHERE dataset=%s ORDER BY id DESC LIMIT 1',('planning:'+area,)).fetchone()
            newer = bool(observed and observed['remote_sha256'] and observed['remote_sha256']!=imported['source_sha256'])
            folder = str((observed or {}).get('remote_filename') or '').rpartition('/')[0]
            path = imported_file(source)
            output.append({'area':area,'snapshot':snap,'imported':imported,'drive':observed,'generation':gen and {k:gen[k] for k in ('id','created_at')},
                           # De onde veio o importado (08/10): 'SAIDA/…' ou só o nome (raiz); pasta None = raiz.
                           'imported_file':path,'imported_folder':(path or '').rpartition('/')[0] or None,
                           'newer_available':newer,
                           # Versão mais recente numa pasta que a carga automática não lê sozinha (nem raiz, nem SAIDA/);
                           # uma versão nova em SAIDA/ entra na carga seguinte e fica «por importar», sem pasta.
                           'newer_folder':folder if newer and folder and folder not in auto_folders() else None})
        return needs.serial({'sources':output,'notice':'Consultar a base não atualiza o Excel nem confirma uma consulta direta ao CPIS.'})
