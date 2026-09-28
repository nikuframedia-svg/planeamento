"""Simplified editor contracts and automatic decisions on disposable sources."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime
from pathlib import Path
import hashlib
import uuid
import zipfile
import pytest
import psycopg
from psycopg.types.json import Jsonb
from app import planning, planning_catalogs as catalogs, planning_needs as needs
from app import planning_associations as assoc, planning_registration as registration, planning_output
from tests.test_planning_needs import canonical, registry, postgres16, vals, create, save, request

ROOT=Path(__file__).resolve().parents[1]


@pytest.mark.parametrize('instant,expected',[
 ('2026-01-28T12:00:00+00:00','2026-02-04'),
 ('2026-12-28T12:00:00+00:00','2027-01-04'),
 ('2026-03-27T23:30:00+00:00','2026-04-03'),
 ('2026-10-24T23:30:00+00:00','2026-11-01'),
 ('2026-06-30T23:30:00+00:00','2026-07-08'),
])
def test_calendar_forecast(instant,expected):
    assert registration.forecast(datetime.fromisoformat(instant),'Europe/Lisbon').isoformat()==expected


def test_forecast_only_after_success_and_stable_across_pieces(canonical):
    n=create()
    assert registration.read('4200') is None
    invalid=vals();invalid['length_mm']=-1
    with pytest.raises(planning.PlanningError):save(n,invalid)
    assert registration.read('OF4200') is None
    p=request(need_id=n['need_id'],expected_revision=n['revision'],area='perfis',catalog_version='s1',values=vals())
    p.pop('actor')
    first=needs.save(p);assert needs.save(p)==first
    stamp=first['registration'];assert stamp==registration.read('4200')
    changed=vals();changed['notes']='Atualizar observação';save(first,changed)
    def another(ref):
        v={**vals(),'component_ref':ref}
        n=needs.resolve(request(area='perfis',production_order_no='4200',values=v))
        return save(n,v)['registration']
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert all(r==stamp for r in pool.map(another,['OTHER-A','OTHER-B']))
    with planning.connect(readonly=True) as conn:
        assert conn.execute('SELECT count(*) AS n FROM planning_mtg.order_registration').fetchone()['n']==1
        assert conn.execute('SELECT min(actor) AS actor FROM planning_mtg.records').fetchone()['actor']=='Utilizador não identificado'


def test_migration_uses_historical_record_creation_and_is_repeatable(canonical):
    n=save(create())
    with psycopg.connect(canonical,autocommit=True) as conn:
        conn.execute('TRUNCATE planning_mtg.order_registration')
        conn.execute("UPDATE planning_mtg.records SET created_at='2025-12-29T22:00:00Z'")
        conn.execute((ROOT/'sql/022_planning_order_registration.sql').read_text())
        conn.execute((ROOT/'sql/022_planning_order_registration.sql').read_text())
    assert registration.read('4200')['predicted_material_request_date']=='2026-01-05'
    assert planning.get_record(n['record_id'])['record']['revision']==1


def direct(canonical):
    from app import cpis_sync
    from tests.test_planning_registry import direct_row
    from datetime import timezone
    cpis_sync.publish([direct_row('OF4200')],datetime.now(timezone.utc),central_dsn=canonical)


def conference(n,quantity):
    proof=assoc.get_evidence(n['need_id'],n['operation_id'])
    p=request(need_id=n['need_id'],operation_id=n['operation_id'],expected_revision=proof['need_revision'],
        evidence_hash=proof['evidence']['evidence_hash'],accepted_required=100,accepted_remaining=quantity,reason='Saldo físico conferido')
    p.pop('actor');return assoc.confer(p)


@pytest.mark.parametrize('balance',[70,0])
def test_full_balance_calculated_only_on_conclusion(canonical,balance):
    n=save(create());direct(canonical)
    assert needs.detail(n['need_id'])['records'][0]['values_json']['quantity_to_plan'] is None
    with pytest.raises(planning.PlanningError,match='quantidade em falta'):save(n,record_status='ready')
    conference(n,balance)
    # Even an old caller supplying a partial quantity cannot override the rule.
    n=save(n,{**vals(),'quantity_to_plan':5},record_status='ready')
    record=needs.detail(n['need_id'])['records'][0]
    assert record['values_json']['quantity_to_plan']==balance
    state=next(f for f in needs.history(n['need_id'])['fields'] if f['field']=='quantity_to_plan')
    assert state['actor']=='Sistema' and state['human_decision'] is None
    if balance==0:assert planning_output._macro_rows([record])==[]
    # Notes-only draft must preserve a previous concluded quantity.
    n=save(n,{**vals(),'notes':'Só observações'})
    assert needs.detail(n['need_id'])['records'][0]['values_json']['quantity_to_plan']==balance


def test_abocardar_checkbox_and_hidden_values_are_preserved(canonical):
    v={**vals(),'abocardar':True,'chanfro':'antigo','ponteira':'histórica','material_request_date':'2025-12-01','picking_week':39,'expected_date':'2025-11-27'}
    n=save(create(),v)
    detail=needs.detail(n['need_id']);assert [o['code'] for o in detail['operations']]==['corte','abocardar']
    assert len(detail['records'])==1 and detail['records'][0]['values_json']['abocardar']=='X'
    n=save(n,{**vals(),'abocardar':False})
    values=needs.detail(n['need_id'])['records'][0]['values_json']
    assert values['abocardar']=='-'
    for field in ('chanfro','ponteira','material_request_date','picking_week','expected_date'):assert values[field]==v[field]
    with planning.connect(readonly=True) as conn:
        assert conn.execute('SELECT count(*) AS n FROM mes_kanban.production_records').fetchone()['n']==0
    n=save(n,{**vals(),'abocardar':'???'})
    assert needs.detail(n['need_id'])['records'][0]['values_json']['abocardar']=='???'
    direct(canonical)
    with pytest.raises(planning.PlanningError) as error:save(n,{**vals(),'abocardar':'???'},record_status='ready')
    assert 'abocardar' in error.value.fields


def test_catalog_families_manual_and_independent_lists(canonical):
    with psycopg.connect(canonical) as conn:
        for index,row in [(4,['Varão quadrado']),(5,['Perfil U']),(6,[None]*17+['UPN50x25']),(7,[None]*17+['UPN50x38'])]:
            conn.execute("INSERT INTO raw_mtg.other_sheet_rows VALUES ('s1','AreaSecaoCorte',%s,%s)",(index,Jsonb({'values':row})))
        header=[None]*22+['Tubo redondo'];header[17]='Perfil U'
        conn.execute("UPDATE raw_mtg.other_sheet_rows SET row_data=%s WHERE sheet_name='AreaSecaoCorte' AND excel_row=1",(Jsonb({'values':header}),))
        conn.execute("INSERT INTO raw_mtg.other_sheet_rows VALUES ('s1','Dados',4,%s)",(Jsonb({'values':[None,None,None,None,'INVENTADO',None,'Equipa sem máquina']}),))
    cat=catalogs.catalog('perfis')
    assert 'INVENTADO' not in cat['material_types'] and 'Equipa sem máquina' in cat['teams']
    assert cat['profiles']['perfil u']==['UPN50x25','UPN50x38']
    assert cat['profile_modes']['varão quadrado']=='manual'
    manual={**vals(),'material_type':'Varão quadrado','profile':'Designação do desenho','width_mm':'20,5'}
    assert catalogs.validate(manual,cat,ready=True)['width_mm']==20.5
    with pytest.raises(planning.PlanningError):catalogs.validate({**manual,'material_type':'Perfil U'},cat,ready=True)
    assert [op['value'] for op in cat['operations']]==['corte']


def test_xml_validation_extensions_and_named_ranges(tmp_path):
    from openpyxl import Workbook
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    book=Workbook();sheet=book.active;sheet.title='Planeamento';book.create_sheet('AreaSecaoCorte')
    book.defined_names.add(DefinedName('PerfilU',attr_text='AreaSecaoCorte!$R$2:$R$150'))
    validation=DataValidation(type='list',formula1='INDIRECT(SUBSTITUTE($AF7," ",""))');sheet.add_data_validation(validation);validation.add('AG7:AG99')
    path=tmp_path/'catalog.xlsx';book.save(path)
    with zipfile.ZipFile(path) as z:entries={n:z.read(n) for n in z.namelist()}
    extra='''<extLst><ext uri="test"><x14:dataValidations xmlns:x14="http://schemas.microsoft.com/office/spreadsheetml/2009/9/main" xmlns:xm="http://schemas.microsoft.com/office/excel/2006/main"><x14:dataValidation type="list"><x14:formula1><xm:f>AreaSecaoCorte!$A$3:$A$700</xm:f></x14:formula1><xm:sqref>AF7:AF99</xm:sqref></x14:dataValidation></x14:dataValidations></ext></extLst>'''
    entries['xl/worksheets/sheet1.xml']=entries['xl/worksheets/sheet1.xml'].replace(b'</worksheet>',extra.encode()+b'</worksheet>')
    with zipfile.ZipFile(path,'w') as z:
        for n,v in entries.items():z.writestr(n,v)
    stat=path.stat();metadata=catalogs.workbook_ranges(str(path),stat.st_mtime_ns,stat.st_size,hashlib.sha256(path.read_bytes()).hexdigest())
    assert metadata['names']['PerfilU']==(17,2,150)
    assert {v['range'] for v in metadata['validations']}=={'AF7:AF99','AG7:AG99'}


def test_old_abocardar_record_cannot_override_checkbox(canonical):
    n=save(create())
    record=needs.detail(n['need_id'])['records'][0]
    with psycopg.connect(canonical) as conn:
        oid=uuid.uuid4()
        conn.execute("INSERT INTO planning_mtg.need_operations VALUES (%s,%s,'perfis','abocardar',2)",(oid,n['need_id']))
        conn.execute("INSERT INTO planning_mtg.records(id,area,production_order_no,component_ref,source_payload,values_json,actor,need_id,operation_id) VALUES (%s,'perfis','OF4200','NEW-A','{}',%s,'Autor histórico',%s,%s)",(uuid.uuid4(),Jsonb({**vals(),'operation':'abocardar'}),n['need_id'],oid))
    with pytest.raises(planning.PlanningError,match='abocardar difere'):planning_output._macro_rows([record])


@pytest.mark.parametrize('checked,expected',[(True,'X'),(False,'-')])
def test_output_checkbox_preserves_removed_cells_and_production(checked,expected,monkeypatch):
    import io
    from xml.etree import ElementTree as ET
    from openpyxl import load_workbook
    from app.dossiers import macro
    from tests.test_dossier_macro import workbook_source, export_piece
    source=workbook_source()
    with zipfile.ZipFile(io.BytesIO(source)) as z:entries={n:z.read(n) for n in z.namelist()}
    namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    ET.register_namespace('',namespace[1:-1])
    sheet=ET.fromstring(entries['xl/worksheets/sheet1.xml'])
    row=sheet.find(f'.//{namespace}row[@r="7"]')
    cell=ET.SubElement(row,namespace+'c',{'r':'U7','t':'inlineStr'});ET.SubElement(ET.SubElement(cell,namespace+'is'),namespace+'t').text='PRESERVAR'
    entries['xl/worksheets/sheet1.xml']=ET.tostring(sheet)
    buffer=io.BytesIO()
    with zipfile.ZipFile(buffer,'w') as z:
        for name,data in entries.items():z.writestr(name,data)
    source=buffer.getvalue();context=export_piece(source)['context']
    context.update(snapshot={'snapshot_id':'s1'},plan_rows=[{'source_line_id':'s1:7','excel_row':7}])
    monkeypatch.setattr(planning_output.cpis,'read_context',lambda of:context)
    record={'id':str(uuid.uuid4()),'production_order_no':'OF265931','component_ref':'REF-A','source_plan_key':'s1:7',
            'actor':'Autor histórico','values_json':{**vals(),'abocardar':catalogs.abocardar_mark(checked),'chanfro':'NÃO ESCREVER','angle_deg':45}}
    rows=planning_output._macro_rows([record]);output,summary=macro.fill_macro(rows,source_bytes=source)
    sheet=load_workbook(io.BytesIO(output),data_only=False)['Planeamento']
    assert sheet['T7'].value==expected and sheet['U7'].value=='PRESERVAR'
    assert sheet['V7'].value==2 and sheet['W7'].value==1
    assert sheet['N7'].value=='=AH7' and sheet['AN7'].value is None
    assert sheet['I7'].value=='2026-09-10' and sheet['AR7'].value=='2026-09-18'
    with zipfile.ZipFile(io.BytesIO(source)) as before,zipfile.ZipFile(io.BytesIO(output)) as after:
        assert before.read('xl/vbaProject.bin')==after.read('xl/vbaProject.bin')
    assert not any(c['column']=='U' for c in summary['cells'])


def test_macro_balance_used_only_for_compatible_piece(canonical):
    source=needs.resolve(request(area='perfis',source={'kind':'plan_line','id':'s1:10','version':'s1'}))
    values={**needs.detail(source['need_id'])['need']['specification'],'operation':'corte','outer_diameter_mm':88.9,'thickness_mm':3,'machine':'MEBA'}
    n=save(source,values);direct(canonical)
    assert assoc.get_evidence(n['need_id'],n['operation_id'])['evidence']['macro_remaining']==36
    n=save(n,values,record_status='ready');assert n['quantity_to_plan']==36
    values['length_mm']=2000;n=save(n,values)
    proof=assoc.get_evidence(n['need_id'],n['operation_id'])
    assert proof['evidence']['macro_remaining'] is None and proof['evidence']['balance_reason']
    with pytest.raises(planning.PlanningError,match='quantidade em falta'):save(n,values,record_status='ready')
    conference(n,30)
    assert save(n,values,record_status='ready')['quantity_to_plan']==30
