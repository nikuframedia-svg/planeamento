"""Closure is a query population, not a deletion or production adjustment."""
import copy
import csv
import io
import uuid

import pytest
import psycopg
from psycopg.rows import dict_row
from openpyxl import load_workbook
from psycopg.types.json import Jsonb
from app import planning, planning_population as population, planning_raw
from app.raw import projection, query, analysis, exports, objects, capacity, edits
from tests.test_raw_workspace import workspace, database, canonical, registry, postgres16


@pytest.mark.parametrize('status,macro,active', [
    ('Em Aberto', '', True), (' Fechada ', '', False),
    ('Em Produção', ' x ', False), ('FECHADA', True, False),
    ('Pronta', '-', True), ('Novo estado', '', True),
    (None, '#VALUE!', True), ('Em Produção', False, True),
])
def test_explicit_closure_normalization(status, macro, active):
    row = {'values': {'status': status, 'final_pct': 125}, 'raw': {'Fechado': macro}}
    assert population.classify(row)['active'] is active
    assert population.includes(row) is active
    assert population.includes(row, 'history') is not active
    assert population.includes(row, 'all')
    if status == 'Novo estado' or macro == '#VALUE!':
        assert population.classify(row)['unknown_states']


def matrix():
    rows = []
    for i, (status, macro) in enumerate([
        ('Em Aberto', ''), ('Fechada', ''), ('Em Aberto', 'X'), ('Fechada', 'X'),
        ('Pronta', ''), ('Desconhecido', ''),
    ]):
        row = {'key': f'piece-{i}', 'values': {
            'of': 'OF4200', 'component_ref': f'PIECE-{i}', 'status': status,
            'machine': 'MEBA', 'operation': 'corte', 'quantity_required': 100,
            'remaining': 10*(i+1), 'expected_date': '2026-09-23',
            'length_mm': 1000, 'section_unit': 10, 'final_pct': 100,
        }, 'raw': {'Fechado': macro, 'Quantidade Prevista': 10*(i+1)},
            'revision': 0, 'need_id': None, 'status_values': [status], 'warnings': []}
        # Closing a macro must also suppress locally prepared work.
        if i == 2:
            row['preparations'] = [{'values_json': {'operation': 'corte', 'quantity_to_plan': 30},
                                    'record_status': 'ready', 'operation_id': 'corte'}]
        rows.append(row)
    return rows


@pytest.mark.parametrize('area', ['perfis', 'cantoneiras'])
def test_population_is_applied_before_pagination_and_facets(workspace, area):
    rows=[]
    expected={scope:set() for scope in ('active','history','all')}
    for batch in range(10):
        for index,row in enumerate(matrix()):
            row['key']=f'{batch:02}-{row["key"]}'
            row['values']['component_ref']=row['key']
            rows.append(row)
            expected['all'].add(row['key'])
            expected['active' if index in (0,4,5) else 'history'].add(row['key'])
    with planning.connect() as c:
        projection.publish(c,'planning:'+area,'pagination-matrix',rows,{})
    for scope,keys in expected.items():
        params={'area':area,'population':scope,'page_size':25,'q':'piece-','state':'all'}
        pages=[query.listing({**params,'page':page}) for page in range(1,5)]
        assert all(page['total']==len(keys) for page in pages)
        observed=[row['key'] for page in pages for row in page['rows']]
        assert set(observed)==keys and len(observed)==len(keys)
        assert len(pages[0]['rows'])==25 and pages[-1]['rows']==[]
        assert sum(option['n'] for option in query.options({**params,'field':'status'})['options'])==len(keys)


@pytest.mark.parametrize('annotated', [False, True])
def test_same_population_in_queries_exports_analysis_capacity_and_old_views(workspace, annotated):
    rows = matrix()
    if annotated:
        for r in rows: population.annotate(r)
    with planning.connect() as c:
        gen = projection.publish(c, 'planning:perfis', 'matrix', rows, {'core_source_fingerprint':projection.fingerprint(c,'perfis')})
        projection.publish(c,'planning:cantoneiras','matrix',[],{'core_source_fingerprint':projection.fingerprint(c,'cantoneiras')})
        projection.publish_orders(c, 'perfis', rows, {}, 'matrix')
        projection.publish(c, 'production:perfis', 'matrix', [
            {'key': 'historical-event', 'planning_keys': ['piece-1'],
             'values': {'of': 'OF4200', 'quantity': 100, 'machine': 'MEBA', 'operation': 'corte'}}], {})
        projection.publish(c, 'production_hours:perfis', 'matrix', [
            {'key': 'historical-hours', 'sheet_uid': 'sheet',
             'values': {'machine': 'MEBA', 'hours_worked': 2, 'production_date': '2026-09-23'}}], {})
    expected = {'active': {'piece-0', 'piece-4', 'piece-5'},
                'history': {'piece-1', 'piece-2', 'piece-3'},
                'all': {r['key'] for r in rows}}
    for scope, keys in expected.items():
        p = {'population': scope, 'page_size': 25}
        result = query.listing(p)
        assert {r['key'] for r in result['rows']} == keys
        assert result['total'] == len(keys)
        assert query.listing({**p, 'q': 'PIECE-', 'page': 2})['rows'] == []
        assert query.listing({**p, 'selected': [r['key'] for r in rows]})['total'] == len(keys)
        assert sum(r['n'] for r in query.options({**p, 'field': 'status'})['options']) == len(keys)
        assert {r['key'] for r in planning_raw.filtered({'rows': rows}, p)} == keys
        exported = list(csv.reader(io.StringIO(exports.table({**p, 'columns': ['component_ref']}, 'csv')[0]), delimiter=';'))
        assert {r[0] for r in exported[1:]} == {r['values']['component_ref'] for r in rows if r['key'] in keys}
        book = load_workbook(io.BytesIO(exports.table({**p, 'columns': ['component_ref']}, 'xlsx')[0]))
        assert book['RAW'].max_row == len(keys)+1
        with planning.connect(readonly=True) as c:
            result = analysis.execute(c, {'area': 'perfis', 'population': scope, 'metrics': [{'name': 'Saldo', 'expression': 'sum([remaining])'}]})
        assert result['rows'] == len(keys)
        assert result['groups'][0]['m0'] == sum(r['values']['remaining'] for r in rows if r['key'] in keys)
        orders = query.listing({**p, 'dataset': 'orders'})['rows']
        assert sum(r['values']['lines_total'] for r in orders) == len(keys)
        assert sum(r['values']['remaining'] for r in orders) == result['groups'][0]['m0']
    old_view = objects.normalize_view({'state': 'all', 'columns': ['of']}, 'perfis')
    assert old_view['population'] == 'active'
    assert query.listing(old_view)['total'] == 3
    assert query.listing({'q': 'PIECE-1'})['total'] == 0
    assert query.listing({'dataset': 'production'})['total'] == 1
    assert query.listing({'dataset': 'production_hours'})['total'] == 1
    capacity.rebuild()
    items = query.listing({'dataset': 'capacity_items'})['rows']
    assert {r['planning_key'] for r in items} == expected['active']
    with planning.connect(readonly=True) as c:
        base, args = query.source(gen)
        assert c.execute('SELECT count(*) n'+base, args).fetchone()['n'] == 6
    with pytest.raises(planning.PlanningError): query.listing({'population': 'something-else'})


@pytest.mark.parametrize('area,source_snapshot,piece_count', [('perfis','s1',2),('cantoneiras','c1',1)])
def test_closing_reopening_sources_retains_identity_and_production(workspace, area, source_snapshot, piece_count):
    # Real projection from source fixtures; each valid source revision is immutable.
    def next_revision(c, suffix, status, closed):
        c.execute("INSERT INTO audit_mtg.snapshots SELECT %s,dataset_id,source_filename,source_path,source_sha256,now()+interval '1 second' FROM audit_mtg.snapshots WHERE snapshot_id=%s", (suffix, source_snapshot))
        for table in ('raw_mtg.plan_production_rows', 'analytics_mtg.kanban_plan_lines', 'core_mtg.production_orders', 'raw_mtg.cpis_rows', 'raw_mtg.other_sheet_rows', 'raw_mtg.machine_rows'):
            schema, name = table.split('.')
            columns = [r['column_name'] for r in c.execute('SELECT column_name FROM information_schema.columns WHERE table_schema=%s AND table_name=%s ORDER BY ordinal_position', (schema, name)).fetchall()]
            expressions = ['%s' if k == 'snapshot_id' else k for k in columns]
            c.execute('INSERT INTO '+table+' SELECT '+','.join(expressions)+" FROM "+table+" WHERE snapshot_id=%s", (suffix, source_snapshot))
        c.execute('UPDATE raw_mtg.plan_production_rows SET closed_x=%s,row_data=row_data||%s WHERE snapshot_id=%s', (closed, Jsonb({'Fechado': 'X' if closed else ''}), suffix))
        c.execute('UPDATE raw_mtg.cpis_rows SET status=%s WHERE snapshot_id=%s', (status, suffix))
    first = projection.rebuild(area)
    initial = query.listing({'area':area})['rows']
    assert len(initial) == piece_count
    with planning.connect(readonly=True) as c:
        evidence_before = c.execute('SELECT * FROM mes_kanban.production_records ORDER BY id').fetchall()
    # The newest CPIS copy (loaded one second later) says Fechada: it decides (decisão de 06/10/2026).
    with psycopg.connect(workspace, row_factory=dict_row) as c: next_revision(c, 'closure', 'Fechada', False)
    closed_gen = projection.rebuild(area)
    assert query.listing({'area':area})['total'] == 0
    history = query.listing({'area':area,'population': 'history'})['rows']
    assert {r['key'] for r in history} == {r['key'] for r in initial}
    assert all(r['population']['closed_sources'] == ['CPIS'] for r in history)
    assert query.listing({'area':area,'version': str(first['id'])})['total'] == piece_count
    with pytest.raises(planning.PlanningError, match='Linha não encontrada'):
        edits.update_batch({'request_id': str(uuid.uuid4()), 'area': area,
            'version': str(closed_gen['id']), 'edits': [{'key': initial[0]['key'],
                'expected_revision': 0, 'values': {'notes': 'Must not enter active batch'}}]})
    with psycopg.connect(workspace, row_factory=dict_row) as c: next_revision(c, 'macroclosed', 'Em Produção', True)
    projection.rebuild(area)
    assert query.listing({'area':area})['total'] == 0
    assert all(r['population']['closed_sources'] == ['Macro'] for r in query.listing({'area':area,'population': 'history'})['rows'])
    with psycopg.connect(workspace, row_factory=dict_row) as c: next_revision(c, 'reopened', 'Em Produção', False)
    projection.rebuild(area)
    assert {r['key'] for r in query.listing({'area':area})['rows']} == {r['key'] for r in initial}
    assert query.listing({'area':area,'population': 'history'})['total'] == 0
    with planning.connect(readonly=True) as c:
        assert c.execute('SELECT * FROM mes_kanban.production_records ORDER BY id').fetchall() == evidence_before
        assert c.execute('SELECT count(*) n FROM raw_mtg.plan_production_rows').fetchone()['n'] == 3 + 3 * piece_count
