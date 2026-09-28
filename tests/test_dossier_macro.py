"""O ficheiro de saída conserva VBA, fórmulas e factos de produção."""
from __future__ import annotations

import hashlib
import io
import re
from xml.etree import ElementTree as ET
from zipfile import ZipFile

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from app.dossiers import DossierError, macro
from tests.test_dossiers import piece_values


def workbook_source():
    book=Workbook();sheet=book.active;sheet.title="Planeamento"
    for col,name in macro.EXPECTED_HEADERS.items():sheet[f"{col}6"]=name
    sheet['A7']=15;sheet['E7']='OF265931';sheet['L7']='REF-A'
    sheet['AF7']='Tubo redondo';sheet['AG7']='114x3';sheet['AH7']=6
    sheet['AM7']=1200;sheet['AO7']='S275JR';sheet['V7']=2;sheet['W7']=1
    sheet['I7']='2026-09-10';sheet['AR7']='2026-09-18';sheet['AS7']=4
    sheet['N7']='=AH7';sheet['Q7']='=AM7';sheet['J7']='=CONCAT(L7," ",AG7)'
    sheet['C7']='=IF(E7="OF265931","Cliente","")'
    sheet['AP7']='=AH7*AI7*AL7';sheet['AT7']='Observação existente'
    sheet['AE7']='=IF(T7="X","X","")'
    for col in ['E','L','AF','AG','AH','AM','AO','V','W','N','J','AP','AT']:
        sheet[f'{col}8'].fill=PatternFill('solid',fgColor='FFFFFF')
    stream=io.BytesIO();book.save(stream)
    out=io.BytesIO()
    with ZipFile(stream) as src,ZipFile(out,'w') as dst:
        for entry in src.infolist():dst.writestr(entry,src.read(entry.filename))
        dst.writestr('xl/vbaProject.bin',b'opaque VBA project bytes preserved without editing')
        dst.writestr('customXml/item1.xml',b'<extension>preserve exactly</extension>')
    return out.getvalue()


def export_piece(source, *, kind='new', changes=None, values=None):
    return {'production_order':'OF265931','values':values or piece_values(),'machine_group':'Serrote',
        'filename':'OF265931.pdf','drawing_pages':[2],'index_page':1,
        'match':{'kind':kind,'excel_row':7,'differences':changes or []},
        'context':{'source':{'source_path':'/unused/Met2_Plan_Perfis.xlsm','source_sha256':hashlib.sha256(source).hexdigest()},
                   'tail':{'last_row':7,'last_id':15},'machines':['Vanguard','MEBA']}}


def output_workbook(data):
    return load_workbook(io.BytesIO(data),data_only=False)['Planeamento']


def test_macro_adds_inputs_preserves_vba_other_parts_and_production():
    source=workbook_source()
    output,summary=macro.fill_macro([export_piece(source)],source_bytes=source)
    assert {k:summary[k] for k in ('added','updated','unchanged')}=={'added':1,'updated':0,'unchanged':0}
    assert {c['column'] for c in summary['cells']} >= {'A','E','L','AF','AG','AH','AM','AO','AT'}
    with ZipFile(io.BytesIO(source)) as original,ZipFile(io.BytesIO(output)) as result:
        assert original.namelist()==result.namelist()
        for name in original.namelist():
            if name not in ('xl/worksheets/sheet1.xml','xl/workbook.xml'):
                assert original.read(name)==result.read(name),name
        assert b'fullCalcOnLoad="1"' in result.read('xl/workbook.xml')
    sheet=output_workbook(output)
    assert sheet['E8'].value=='OF265931' and sheet['AH8'].value==6 and sheet['AM8'].value==1200
    assert sheet['AF8'].value=='Tubo redondo' and sheet['A8'].value==16
    assert sheet['N8'].value=='=AH8' and sheet['J8'].value=='=CONCAT(L8," ",AG8)'
    assert sheet['V8'].value is None and sheet['W8'].value is None
    assert sheet['V7'].value==2 and sheet['W7'].value==1
    assert sheet['I7'].value=='2026-09-10' and sheet['AR7'].value=='2026-09-18'
    assert sheet['AQ8'].value is None  # «Serrote» não escolhe um equipamento.
    assert 'OF265931.pdf' in sheet['AT8'].value


def test_quantity_revision_preserves_completed_quantities_and_formulas():
    source=workbook_source()
    piece=export_piece(source,kind='changed',changes=[{'field':'quantity_required','pdf':10,'plan':6}],values=piece_values(quantity=10))
    output,summary=macro.fill_macro([piece],source_bytes=source)
    sheet=output_workbook(output)
    assert {k:summary[k] for k in ('added','updated','unchanged')}=={'added':0,'updated':1,'unchanged':0}
    assert sheet['AH7'].value==10 and sheet['N7'].value=='=AH7'
    assert sheet['V7'].value==2 and sheet['W7'].value==1
    assert sheet['AS7'].value==4 and sheet['AR7'].value=='2026-09-18'
    assert sheet['AT7'].value.startswith('Observação existente\n')


def test_existing_lines_receive_idempotent_provenance_without_touching_production():
    source=workbook_source()
    output,summary=macro.fill_macro([export_piece(source,kind='existing')],source_bytes=source)
    assert summary['updated']==1 and summary['added']==0
    sheet=output_workbook(output)
    assert sheet['V7'].value==2 and sheet['W7'].value==1
    assert sheet['AT7'].value.startswith('Observação existente\n[dossier:')


def test_formulas_cannot_be_injected_by_a_drawing_reference_or_notes():
    source=workbook_source()
    piece=export_piece(source,values=piece_values(ref='=HYPERLINK("https://bad.example")',notes='=cmd|anything'))
    output,_=macro.fill_macro([piece],source_bytes=source)
    sheet=output_workbook(output)
    assert sheet['L8'].data_type=='s' and sheet['L8'].value.startswith('=HYPERLINK')
    piece['context'].update(sales_order='OV1',customer='=evil',designation='x')
    csv=macro.csv_bytes([piece]).decode('utf-8-sig')
    assert "'=evil" in csv and "'=HYPERLINK" in csv


def test_macro_hash_header_and_occupied_target_are_verified():
    source=workbook_source()
    piece=export_piece(source)
    with pytest.raises(DossierError,match='mudou'):macro.fill_macro([piece],source_bytes=source+b'changed')
    def alter_xml(before,after):
        out=io.BytesIO()
        with ZipFile(io.BytesIO(source)) as src,ZipFile(out,'w') as dst:
            for entry in src.infolist():
                value=src.read(entry.filename)
                if entry.filename=='xl/worksheets/sheet1.xml':value=re.sub(before,after,value) if isinstance(before,re.Pattern) else value.replace(before,after)
                dst.writestr(entry,value)
        return out.getvalue()
    other=alter_xml(b'QTD [un,]',b'Unknown heading')
    with pytest.raises(DossierError,match='colunas'):macro.fill_macro([export_piece(other)],source_bytes=other)
    other=alter_xml(re.compile(rb'<c\b[^>]*\br="E8"[^>]*(?:/>|>.*?</c>)'),b'<c r="E8" t="inlineStr"><is><t>OF999999</t></is></c>')
    # openpyxl pode variar a representação de uma célula vazia; verificar a mutação.
    with ZipFile(io.BytesIO(other)) as z:
        assert b'OF999999' in z.read('xl/worksheets/sheet1.xml')
    with pytest.raises(DossierError,match='já contém'):macro.fill_macro([export_piece(other)],source_bytes=other)


def test_vanguard_assignment_uses_existing_machine_catalogue():
    source=workbook_source();piece=export_piece(source);piece['machine_group']='Vanguard'
    data,_=macro.fill_macro([piece],source_bytes=source)
    assert output_workbook(data)['AQ8'].value=='Vanguard'


def test_existing_vanguard_line_gets_machine_and_angle_is_never_written():
    source=workbook_source();piece=export_piece(source,kind='existing',values=piece_values(angle_deg=68))
    piece['machine_group']='Vanguard';piece['match']['machine_update']='Vanguard'
    data,summary=macro.fill_macro([piece],source_bytes=source);sheet=output_workbook(data)
    assert sheet['AQ7'].value=='Vanguard' and sheet['AN7'].value is None
    assert any(c['column']=='AQ' and c['before'] is None for c in summary['cells'])
    assert 'sem conversão para Ang,' in sheet['AT7'].value


def test_resolved_catalogue_values_are_exported_without_replacing_pdf_literal():
    source=workbook_source();values=piece_values(profile='UPN80',material_type='Perfil U')
    piece=export_piece(source,values=values)
    piece['match'].update(effective_values={**values,'profile':'UPN80x45','material_type':'Perfil U'})
    data,_=macro.fill_macro([piece],source_bytes=source);sheet=output_workbook(data)
    assert sheet['AG8'].value=='UPN80x45' and piece['values']['profile']=='UPN80'


def test_chamfer_and_angle_require_separate_export_decisions():
    source=workbook_source();values=piece_values(chanfro='X',angle_deg=45)
    piece=export_piece(source,values=values)
    data,_=macro.fill_macro([piece],source_bytes=source);sheet=output_workbook(data)
    assert sheet['U8'].value is None and sheet['AN8'].value is None
    piece['match']['chanfro_export']=True
    data,_=macro.fill_macro([piece],source_bytes=source);sheet=output_workbook(data)
    assert sheet['U8'].value=='X' and sheet['AN8'].value is None


def test_dossier_note_is_not_appended_twice_after_output_is_reimported():
    source=workbook_source();piece=export_piece(source,kind='existing')
    first,_=macro.fill_macro([piece],source_bytes=source)
    piece['context']['source']['source_sha256']=hashlib.sha256(first).hexdigest()
    second,summary=macro.fill_macro([piece],source_bytes=first)
    assert summary['unchanged']==1 and summary['updated']==0
    assert output_workbook(second)['AT7'].value.count('[dossier:')==1


def test_dossier_note_updates_its_logical_block_and_preserves_human_text():
    source=workbook_source();first_piece=export_piece(source,kind='existing',values=piece_values(notes='Nota antiga'))
    first,_=macro.fill_macro([first_piece],source_bytes=source)
    second_piece=export_piece(first,kind='existing',values=piece_values(notes='Nota nova\nem duas linhas'))
    second,_=macro.fill_macro([second_piece],source_bytes=first)
    note=output_workbook(second)['AT7'].value
    assert note.startswith('Observação existente\n') and note.count('[dossier:')==1
    assert 'Nota nova em duas linhas' in note and 'Nota antiga' not in note


def test_shared_formula_extraction_does_not_consume_following_cells():
    xml=b'<c r="N7"><f t="shared" si="1" ref="N7:N9">AH7</f><v>6</v></c><c r="N8"><f t="shared" si="1"/><v>0</v></c><c r="Q8"><f>AM8</f></c>'
    assert macro._shared_formulas(xml)=={'1':('N7','AH7')}
